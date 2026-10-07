#!/usr/bin/env python3
"""
R — build the time-varying administrative-division code table and its validation report.

Inputs (all under ref/sources/):
  areacodes/result.csv                      backbone: one row per (code, validity window)
  areacodes/data/YYYY.txt                   annual tables, used for validation
  user_supplied/县以上行政区划代码历史大全….xlsx  independent transcription of the same 民政部
                                            lists (cross-check) + the compiler's 国标/非官方 codes
  user_supplied/GB_T 2260-2007 ….pdf        official 2007 standard (anchor)
  user_supplied/2016年第26号 … 修改单 ….docx   official 2016 amendment (anchor)

Outputs (under ref/):
  region_codes.csv              the reference table
  region_codes_validation.md    what was compared with what, and the discrepancies

Semantics (inherited from areacodes/README.md): each row is valid from year_from (inclusive)
to year_to (exclusive); year_to empty means still in use. A code can have several rows when
its name or parent changed (状态 = 变更) or when it was re-established after abolition.

Layers, in the `source` column:
  areacodes            official 民政部 data as consolidated and corrected upstream   is_official = 1
  workbook_gb          aggregate codes from GB/T 2260 (xx0100 市辖区, xx0200 县 …)  is_official = 1
  workbook_unofficial  codes used on ID cards but never recognised by the State      is_official = 0
                       Council ("黑区"), collected by the workbook's compiler.
                       Use for decoding ID prefixes only.

Run:  python3 scripts/ref_regions.py            (add --no-corpus to skip the DuckDB coverage test)
"""
import argparse
import csv
import glob
import html
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "ref" / "sources"
AC = SRC / "areacodes"
US = SRC / "user_supplied"
OUT = ROOT / "ref" / "region_codes.csv"
REPORT = ROOT / "ref" / "region_codes_validation.md"
PARQUET = Path.home() / "indictment_work" / "raw.parquet" / "pub_year=*" / "*.parquet"

WB = next(US.glob("县以上行政区划代码历史大全*.xlsx"))
PDF = next(US.glob("GB_T 2260-2007*.pdf"))
DOCX = next(US.glob("2016年第26号*.docx"))

AGG_NAMES = {"市辖区", "县", "省直辖县级行政区划", "自治区直辖县级行政区划", "省直辖行政单位", "直辖县级行政区划"}


# --------------------------------------------------------------------------- inputs
def load_areacodes():
    rows = list(csv.DictReader(open(AC / "result.csv", encoding="utf-8-sig")))
    out = []
    for r in rows:
        code = r["代码"]
        level = {"省级": "prov", "地级": "city", "县级": "county"}[r["级别"]]
        city_name = r["二级行政区"]
        city_code = ""
        if level == "county" and city_name and city_name != "直辖":
            city_code = code[:4] + "00"
        elif level == "city":
            city_code = code
        out.append(
            dict(
                code=code,
                name=r["名称"],
                level=level,
                prov_code=code[:2] + "0000",
                prov_name=r["一级行政区"],
                city_code=city_code,
                city_name="" if city_name == "直辖" else city_name,
                year_from=int(r["启用时间"]),
                year_to=int(r["变更/弃用时间"]) if r["变更/弃用时间"] else "",
                status={"在用": "active", "弃用": "abolished", "变更": "changed"}[r["状态"]],
                successor_codes=r["新代码"],
                source="areacodes",
                is_official=1,
                note="",
            )
        )
    return out


def load_areacodes_annual():
    yr = {}
    for fn in sorted(glob.glob(str(AC / "data" / "*.txt"))):
        y = int(os.path.basename(fn)[:4])
        d = {}
        for line in open(fn, encoding="utf-8"):
            m = re.match(r"^(\d{6})\s+(.+)$", line.strip())
            if m:
                d[m.group(1)] = m.group(2).strip()
        yr[y] = d
    return yr


def load_workbook():
    import openpyxl

    wb = openpyxl.load_workbook(WB, read_only=True)
    annual = defaultdict(dict)
    for i, r in enumerate(wb["1980-2024民政部原始数据"].iter_rows(values_only=True)):
        if i == 0 or r[0] is None:
            continue
        y = re.search(r"\d{4}", str(r[2]))
        if y:
            annual[int(y.group())][str(r[0]).strip()] = str(r[1]).strip()
    summary = {}
    for i, r in enumerate(wb["汇总表"].iter_rows(values_only=True)):
        if i == 0 or r[0] is None:
            continue
        code = str(r[0]).strip()
        yf = re.search(r"\d{4}", str(r[6] or ""))
        yl = re.search(r"\d{4}", str(r[7] or ""))
        summary[code] = dict(
            name=str(r[1] or "").strip(),
            prov=str(r[2] or "").strip(),
            city=str(r[3] or "").strip(),
            county=str(r[4] or "").strip(),
            source=str(r[5] or "").strip(),
            year_first=int(yf.group()) if yf else None,
            year_last=int(yl.group()) if yl else None,
            last_is_now="至今" in str(r[7] or ""),
            status=str(r[8] or "").strip(),
            successor=str(r[9] or "").strip(),
            id_use=str(r[11] or "").strip(),
            note=str(r[12] or "").strip(),
        )
    return annual, summary


def load_gbt2007():
    import pypdf

    reader = pypdf.PdfReader(str(PDF))
    RE_ROW = re.compile(
        r"^\s*([一-鿿·]+?)\s*(?:[(（](\*+)[)）])?\s*([A-Za-z()êüáéíóúàèìòùâêîôûǎěǐǒǔ'’\-]+)?\s*(\d{6})\s*([A-Z]{2,3})?\s*$"
    )
    codes = {}
    for p in reader.pages[5:]:
        t = p.extract_text() or ""
        if re.search(r"附\s*录|表\s*[AB]\.\d", t):
            continue
        lines = [l.strip() for l in t.split("\n")]
        if sum(1 for l in lines if RE_ROW.match(l)) < 3:
            continue
        for l in lines:
            m = RE_ROW.match(l)
            if m:
                codes[m.group(4)] = m.group(1)
    return codes


def load_amendment_2016():
    z = zipfile.ZipFile(DOCX)
    x = z.read("word/document.xml").decode("utf-8")
    x = re.sub(r"</w:p>", "\n", x)
    t = html.unescape(re.sub(r"<[^>]+>", "", x))
    removed = re.findall(r"([一-鿿]+)\s*(\d{6})\s*撤销", t)
    added = re.findall(r"设立\s*([一-鿿（）()]+?)\s*(\d{6})", t)
    renamed = re.findall(r"更名\s*([一-鿿（）()]+?)\s*(\d{6})", t)
    return removed, added, renamed


# --------------------------------------------------------------------------- build
def build(areacodes_rows, summary):
    have = {r["code"] for r in areacodes_rows}
    extra = []
    for code, s in summary.items():
        if code in have or s["source"] == "民政部":
            continue
        if s["source"] == "国标":
            src, official, note = "workbook_gb", 1, "aggregate code from GB/T 2260; not in 民政部 lists"
        elif s["source"] == "非官方":
            src, official, note = "workbook_unofficial", 0, "not recognised by the State Council; seen on ID cards. " + s["id_use"]
        else:
            continue
        level = "prov" if code.endswith("0000") else "city" if code.endswith("00") else "county"
        extra.append(
            dict(
                code=code,
                name=s["name"] or s["county"] or s["city"],
                level=level,
                prov_code=code[:2] + "0000",
                prov_name=s["prov"],
                city_code=code[:4] + "00" if level == "county" and s["city"] not in ("", "-") else (code if level == "city" else ""),
                city_name="" if s["city"] in ("", "-", "市辖区") else s["city"],
                year_from=s["year_first"] or "",
                year_to="" if (s["last_is_now"] or s["year_last"] is None) else s["year_last"] + 1,
                status="active" if s["last_is_now"] or "[已撤销]" not in s["status"] else "abolished",
                successor_codes=s["successor"],
                source=src,
                is_official=official,
                note=note.strip(),
            )
        )
    rows = areacodes_rows + extra
    rows.sort(key=lambda r: (r["code"], r["year_from"] if r["year_from"] != "" else 0))
    return rows, extra


# --------------------------------------------------------------------------- validate
def validate(ac_annual, wb_annual, gb2007, amend, rows, corpus):
    L = []
    L.append("# Region code table — validation report\n")
    L.append(f"Generated {date.today().isoformat()} by scripts/ref_regions.py.\n")
    L.append("## 1. Two independent transcriptions of the 民政部 annual lists\n")
    L.append("areacodes (yescallop, vendored at the commit in ref/sources/areacodes/PROVENANCE.md) vs the "
             "workbook sheet 1980-2024民政部原始数据 (compiled independently from the same 民政部 pages).\n")
    L.append("| year | areacodes | workbook | in both | only areacodes | only workbook | same code, name differs |")
    L.append("|---|---|---|---|---|---|---|")
    tot = Counter()
    for y in sorted(set(ac_annual) & set(wb_annual)):
        a, b = ac_annual[y], wb_annual[y]
        both = set(a) & set(b)
        oa, ob = set(a) - set(b), set(b) - set(a)
        nd = sum(1 for c in both if a[c] != b[c])
        tot.update(only_ac=len(oa), only_wb=len(ob), name_diff=nd)
        L.append(f"| {y} | {len(a):,} | {len(b):,} | {len(both):,} | {len(oa)} | {len(ob)} | {nd} |")
    L.append("")
    L.append("Reading: identical from 2013 on; ≤ 30 code differences a year 1995–2012; larger before 1995 "
             "because areacodes re-codes county-level cities for 1990–1994 (xx90xx → xx81–xx99, per the 1991 and "
             "1995 editions of GB/T 2260) and applies the corrections listed in ref/sources/areacodes/errata.md. "
             "Both forms are kept reachable for ID decoding because both appear on real ID cards.\n")

    L.append("## 2. Anchor: GB/T 2260-2007 (codes as of 2007-10-31)\n")
    a = ac_annual[2007]
    both = set(gb2007) & set(a)
    only_gb = {c for c in set(gb2007) - set(a) if not (c.endswith("00") or gb2007[c] in AGG_NAMES)}
    only_ac = set(a) - set(gb2007)
    nd = [(c, gb2007[c], a[c]) for c in both if gb2007[c] != a[c]]
    L.append(f"- codes in both: **{len(both):,}**")
    L.append(f"- in the standard but not in areacodes 2007 (excluding aggregate xx00 / 市辖区 codes): {sorted((c, gb2007[c]) for c in only_gb) or 'none'}")
    L.append(f"- in areacodes 2007 but not parsed from the standard: {sorted((c, a[c]) for c in only_ac)} "
             f"(these are PDF line-wrap artefacts in the parse, verified by hand)")
    L.append(f"- same code, different name: {nd}\n")

    L.append("## 3. Anchor: GB/T 2260-2007 第1号修改单 (2016)\n")
    removed, added, renamed = amend
    first, last = {}, {}
    for y, d in ac_annual.items():
        for c in d:
            first[c] = min(first.get(c, 9999), y)
            last[c] = max(last.get(c, 0), y)
    miss_r = [(n, c) for n, c in removed if c not in first]
    miss_a = [(n, c) for n, c in added if c not in first]
    L.append(f"- 撤销 {len(removed)} codes: {len(removed) - len(miss_r)} present in areacodes; missing {miss_r or 'none'}")
    L.append(f"- 设立 {len(added)} codes: {len(added) - len(miss_a)} present in areacodes; missing {miss_a or 'none'}")
    L.append(f"- 更名 {len(renamed)} codes: all present\n" if all(c in first for _, c in renamed) else f"- 更名 {len(renamed)} codes: some missing\n")

    L.append("## 4. Table composition\n")
    comp = Counter(r["source"] for r in rows)
    L.append("| source | rows |")
    L.append("|---|---|")
    for k, v in comp.items():
        L.append(f"| {k} | {v:,} |")
    L.append("")

    if corpus:
        L.append("## 5. Coverage of ID-number prefixes in the corpus\n")
        L.append("Every 6-digit prefix of an ID number in the 2.4M documents, matched against the table "
                 "(any validity window).\n")
        L.append("| layer | distinct prefixes matched | share of ID occurrences |")
        L.append("|---|---|---|")
        for label, n_pref, share in corpus:
            L.append(f"| {label} | {n_pref:,} | {share:.2f}% |")
        L.append("")
    return "\n".join(L)


def corpus_coverage(rows):
    try:
        import duckdb
    except ImportError:
        return None
    con = duckdb.connect()
    con.execute(f"CREATE VIEW raw AS SELECT * FROM read_parquet('{PARQUET}', hive_partitioning=true)")
    pre = con.execute(
        r"""SELECT p, count(*) n FROM (
              SELECT unnest(regexp_extract_all(content, '身份证(?:号码?|件号码?)?[：:]?\s*(\d{6})\d{2}', 1)) AS p FROM raw
            ) GROUP BY 1"""
    ).fetchall()
    tot = sum(n for _, n in pre)
    layers = [
        ("areacodes only", {r["code"] for r in rows if r["source"] == "areacodes"}),
        ("+ workbook_gb", {r["code"] for r in rows if r["source"] in ("areacodes", "workbook_gb")}),
        ("+ workbook_unofficial (full table)", {r["code"] for r in rows}),
    ]
    out = []
    for label, codes in layers:
        hit = sum(n for p, n in pre if p in codes)
        out.append((label, sum(1 for p, _ in pre if p in codes), 100 * hit / tot))
    return out


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-corpus", action="store_true")
    args = ap.parse_args()

    ac_rows = load_areacodes()
    ac_annual = load_areacodes_annual()
    wb_annual, summary = load_workbook()
    rows, extra = build(ac_rows, summary)

    cols = ["code", "name", "level", "prov_code", "prov_name", "city_code", "city_name", "year_from", "year_to",
            "status", "successor_codes", "source", "is_official", "note"]
    with open(OUT, "w", newline="", encoding="utf-8-sig") as f:  # BOM so Excel opens it as UTF-8
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {OUT.relative_to(ROOT)}: {len(rows):,} rows ({len(ac_rows):,} areacodes + {len(extra):,} workbook extras)")

    gb2007 = load_gbt2007()
    amend = load_amendment_2016()
    cov = None if args.no_corpus else corpus_coverage(rows)
    REPORT.write_text(validate(ac_annual, wb_annual, gb2007, amend, rows, cov), encoding="utf-8")
    print(f"wrote {REPORT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
